"""ACI133 external assessment of GEM v2026 Turkey functions.

The workbook is treated as immutable. Taxonomy attributes absent from the source are
represented by named sensitivity scenarios; they are never written back as observations.
"""
from __future__ import annotations

import hashlib
import json
import math
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
import pandas as pd
from scipy.stats import norm, spearmanr

from fragility_tools import score_states, state_probabilities
from pipeline import evaluate, read_grid, read_model, sample_grid
from rupture_distance import trace_distance

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "data/raw/aci133/ACI133v3.xlsx"
RAW = ROOT / "data/raw"
OUT = ROOT / "outputs/aci133_exploratory"
AUDIT = ROOT / "outputs/aci133_audit"

EVENTS = [
    ("event1", "2023-02-06_M7.8"),
    ("event3", "2023-02-06_M6.7"),
    ("event2", "2023-02-06_M7.5"),
    ("event4", "2023-02-20_M6.3"),
]
EVENT_LABELS = dict(EVENTS)
EVENT_SCOPES = {
    "two_main_events": ["event1", "event2"],
    "four_survey_preceding_events": ["event1", "event3", "event2", "event4"],
}
IM_FIELDS = {"PGA": "PGA", "SA(0.3)": "PSA03", "SA(0.6)": "PSA06", "SA(1.0)": "PSA10"}
DLR = {0: 0.0, 1: 0.05, 2: 0.15, 3: 0.60, 4: 1.0}
DAMAGE_TO_DS = {"N": 0, "L": 1, "M": 2, "S": 3, "C": 4}

DESIGN_MAPS = {
    "baseline": {
        "1975_1997": "CDL+ERL", "1998_2006": "CDM+ERL",
        "2007_2018": "CDM+ERM", "2019_plus": "CDH+ERH"},
    "lower": {
        "1975_1997": "CDN+ERN", "1998_2006": "CDL+ERL",
        "2007_2018": "CDM+ERL", "2019_plus": "CDM+ERM"},
    "upper": {
        "1975_1997": "CDL+ERM", "1998_2006": "CDM+ERM",
        "2007_2018": "CDH+ERM", "2019_plus": "CDH+ERH"},
}
SYSTEMS = {"lfinf": "LFINF", "lfm": "LFM", "lwal": "LWAL"}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write_csv(frame: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    frame.to_csv(path, index=False, encoding="utf-8-sig")


def unique_headers(values) -> list[str]:
    counts: dict[str, int] = {}
    result = []
    for value in values:
        base = str(value).strip() if value is not None else "unnamed"
        counts[base] = counts.get(base, 0) + 1
        result.append(base if counts[base] == 1 else f"{base}__{counts[base]}")
    return result


def load_workbook_records() -> tuple[pd.DataFrame, pd.DataFrame]:
    book = openpyxl.load_workbook(SOURCE, data_only=True, read_only=False)
    sheet = book[book.sheetnames[0]]
    headers = unique_headers([sheet.cell(2, c).value for c in range(1, sheet.max_column + 1)])
    units = [sheet.cell(3, c).value for c in range(1, sheet.max_column + 1)]
    records = []
    for r in range(4, sheet.max_row + 1):
        row = {headers[c-1]: sheet.cell(r, c).value for c in range(1, sheet.max_column + 1)}
        red_fields = []
        for c, header in enumerate(headers, 1):
            color = sheet.cell(r, c).font.color
            if color and color.type == "rgb" and color.rgb == "FFFF0000":
                red_fields.append(header)
        row["source_supplemented_fields"] = "|".join(red_fields)
        row["source_supplemented_count"] = len(red_fields)
        records.append(row)
    dictionary = pd.DataFrame({"source_column": headers, "unit_as_supplied": units})
    return pd.DataFrame(records), dictionary


def era_from_age(age) -> str | None:
    if pd.isna(age):
        return None
    estimated_year = 2023 - float(age)
    if estimated_year < 1975:
        return "pre1975"
    if estimated_year <= 1997:
        return "1975_1997"
    if estimated_year <= 2006:
        return "1998_2006"
    if estimated_year <= 2018:
        return "2007_2018"
    return "2019_plus"


def prepare() -> pd.DataFrame:
    source, dictionary = load_workbook_records()
    selected = pd.DataFrame({
        "building_id": source["ID"].astype(str),
        "age_years_reported": pd.to_numeric(source["Age"], errors="coerce"),
        "storeys_reported": pd.to_numeric(source["No Stories"], errors="coerce"),
        "height_m": pd.to_numeric(source["Total Height"], errors="coerce"),
        "floor_area_m2": pd.to_numeric(source["Floor Area"], errors="coerce"),
        "structural_damage_code": source["Structural Damage (5-class)"].astype(str),
        "critical_damage_flag": source["Critical Damage"].astype(str),
        "soft_story_flag": source["Soft Story"].astype(str),
        "latitude": pd.to_numeric(source["Latitude"], errors="coerce"),
        "longitude": pd.to_numeric(source["Longitude"], errors="coerce"),
        "vs30_source_mps": pd.to_numeric(source["V_S30"], errors="coerce"),
        "concrete_wall_area_ns_m2": pd.to_numeric(source["C.F. Concrete Wall Area NS"], errors="coerce"),
        "concrete_wall_area_ew_m2": pd.to_numeric(source["C.F. Concrete Wall Area EW"], errors="coerce"),
        "min_wall_index_as_supplied": pd.to_numeric(source["min WI (%)"], errors="coerce"),
        "source_supplemented_fields": source["source_supplemented_fields"],
        "source_supplemented_count": source["source_supplemented_count"],
    })
    selected["construction_year_estimated"] = 2023 - selected.age_years_reported
    selected["code_era_proxy"] = selected.age_years_reported.map(era_from_age)
    selected["observed_ds"] = selected.structural_damage_code.map(DAMAGE_TO_DS)
    selected["observed_loss_proxy"] = selected.observed_ds.map(DLR)
    selected["coordinate_duplicate_group_size"] = selected.groupby(
        ["latitude", "longitude"])["building_id"].transform("size")
    selected["has_any_concrete_wall"] = selected[["concrete_wall_area_ns_m2",
                                                    "concrete_wall_area_ew_m2"]].max(axis=1) > 0
    main_trace_distances = np.column_stack([
        trace_distance(RAW / "event1_rupture.json", selected.latitude, selected.longitude),
        trace_distance(RAW / "event2_rupture.json", selected.latitude, selected.longitude),
    ])
    selected["min_mainshock_trace_distance_km"] = np.nanmin(main_trace_distances, axis=1)
    if selected.building_id.duplicated().any():
        raise ValueError("Duplicate building IDs in ACI133 workbook")
    if not selected.structural_damage_code.isin(DAMAGE_TO_DS).all():
        raise ValueError("Unrecognised ACI structural damage code")
    write_csv(dictionary, AUDIT / "source_dictionary.csv")
    missing = selected.isna().mean().rename("missing_fraction").reset_index()
    missing.columns = ["field", "missing_fraction"]
    write_csv(missing, AUDIT / "missingness.csv")
    write_csv(selected.structural_damage_code.value_counts().rename_axis("damage_code").reset_index(name="n"),
              AUDIT / "damage_counts.csv")
    write_csv(selected, ROOT / "data/processed/aci133_buildings.csv")
    audit = {
        "records": int(len(selected)),
        "unique_ids": int(selected.building_id.nunique()),
        "unique_coordinate_pairs": int(selected[["latitude", "longitude"]].drop_duplicates().shape[0]),
        "source_supplemented_cells": int(selected.source_supplemented_count.sum()),
        "source_supplemented_buildings": int((selected.source_supplemented_count > 0).sum()),
        "exact_gem_height_records": int(selected.storeys_reported.isin([1, 2, 3, 4, 5, 6, 8, 10]).sum()),
        "selection_warning": "Accessibility/size/age/complexity sample; unsafe-to-enter buildings were excluded.",
        "age_interpretation": "Construction year is derived as 2023 minus reported age and is a proxy, not a permit year.",
        "wall_index_unit_warning": "Header says percent, while stored values are fractional-scale; no conversion is imposed here.",
    }
    AUDIT.mkdir(parents=True, exist_ok=True)
    (AUDIT / "audit_summary.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
    return selected


def assign_ims(buildings: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for key, label in EVENTS:
        grid, units, _ = read_grid(RAW / f"{key}_grid.xml")
        for imt, field in IM_FIELDS.items():
            derivation = "direct_shakemap_grid"
            if field in units and units[field] in ("%g", "g"):
                values = sample_grid(grid, field, buildings.latitude, buildings.longitude)
                if units[field] == "%g":
                    values = values / 100.0
            elif imt == "SA(0.6)" and units.get("PSA03") in ("%g", "g") and units.get("PSA10") in ("%g", "g"):
                sa03 = sample_grid(grid, "PSA03", buildings.latitude, buildings.longitude)
                sa10 = sample_grid(grid, "PSA10", buildings.latitude, buildings.longitude)
                if units["PSA03"] == "%g": sa03 = sa03 / 100.0
                if units["PSA10"] == "%g": sa10 = sa10 / 100.0
                weight = math.log(0.6 / 0.3) / math.log(1.0 / 0.3)
                values = np.exp(np.log(sa03) + weight * (np.log(sa10) - np.log(sa03)))
                derivation = "log_period_log_sa_interpolation_from_SA0.3_SA1.0"
            else:
                values = np.full(len(buildings), np.nan)
                derivation = "not_available"
            rows.append(pd.DataFrame({"building_id": buildings.building_id, "event_key": key,
                                      "event_label": label, "imt": imt, "im_g": values,
                                      "im_derivation": derivation,
                                      "source_sha256": sha256(RAW / f"{key}_grid.xml")}))
    result = pd.concat(rows, ignore_index=True)
    write_csv(result, ROOT / "data/processed/aci133_im_assignments.csv")
    return result


def load_fragilities() -> dict[tuple[str, str], dict[str, np.ndarray]]:
    result = {}
    for imt in IM_FIELDS:
        path = RAW / "gem_fragility" / f"structural_fragility_params_{imt}.csv"
        frame = pd.read_csv(path)
        required = [f"Median_DS{i}" for i in range(1, 5)] + [f"Beta_DS{i}" for i in range(1, 5)]
        if not {"Taxonomy", *required}.issubset(frame.columns):
            raise ValueError(f"Unexpected fragility schema: {path}")
        for _, row in frame.iterrows():
            result[(row.Taxonomy, imt)] = {
                "median": row[[f"Median_DS{i}" for i in range(1, 5)]].to_numpy(float),
                "beta": row[[f"Beta_DS{i}" for i in range(1, 5)]].to_numpy(float),
            }
    return result


def fragility_exceedance(parameters, im_values) -> tuple[np.ndarray, np.ndarray]:
    im_values = np.asarray(im_values, float)
    med = np.asarray(parameters["median"], float)
    beta = np.asarray(parameters["beta"], float)
    out = np.full((len(im_values), 4), np.nan)
    valid = np.isfinite(im_values) & (im_values > 0)
    out[valid] = norm.cdf((np.log(im_values[valid, None]) - np.log(med[None, :])) / beta[None, :])
    # The published DS fits can cross. Enforce the probability ordering needed
    # for state probabilities and retain the size of this post-processing step.
    adjusted = np.minimum.accumulate(out, axis=1)
    max_adjustment = np.nanmax(np.abs(adjusted-out), axis=1)
    return adjusted, max_adjustment


def choose_height(storeys: float, allowed: list[int], policy: str) -> tuple[int | None, float | None]:
    if not np.isfinite(storeys) or not allowed:
        return None, None
    if policy == "exact":
        if float(storeys).is_integer() and int(storeys) in allowed:
            return int(storeys), 0.0
        return None, None
    if policy == "nearest":
        value = min(allowed, key=lambda h: (abs(h - storeys), h))
        return value, float(value - storeys)
    raise ValueError(policy)


def vulnerability_metrics(y, pred) -> dict:
    y, pred = np.asarray(y, float), np.asarray(pred, float)
    ok = np.isfinite(y) & np.isfinite(pred)
    y, pred = y[ok], pred[ok]
    rho = spearmanr(y, pred).statistic if len(y) > 2 and len(set(y)) > 1 and len(set(pred)) > 1 else np.nan
    return {"n": len(y), "bias_pred_obs": pred.sum() / y.sum() if y.sum() else np.nan,
            "rmse": np.sqrt(np.mean((pred-y)**2)) if len(y) else np.nan, "spearman": rho}


def analyse(buildings: pd.DataFrame, ims: pd.DataFrame) -> None:
    model = read_model(RAW / "gem_vulnerability_structural.xml")
    fragility = load_fragilities()
    # Structural curves are asserted identical across occupancy for a shared base taxonomy.
    occupancy_check = {}
    for taxonomy, fun in model.items():
        base = taxonomy.rsplit("/", 1)[0]
        signature = (fun["imt"], tuple(fun["x"]), tuple(fun["mean"]), tuple(fun["cov"]))
        occupancy_check.setdefault(base, set()).add(signature)
    if any(len(v) != 1 for v in occupancy_check.values()):
        raise ValueError("Structural vulnerability varies by occupancy; missing ACI occupancy cannot be neutralised")

    im_lookup = ims.pivot_table(index="building_id", columns=["event_key", "imt"], values="im_g", aggfunc="first")
    results = []
    systems = SYSTEMS.items()
    for system_name, system in systems:
        for design_name, design_map in DESIGN_MAPS.items():
            for height_policy in ("exact", "nearest"):
                scenario = f"{system_name}__{design_name}__{height_policy}"
                for _, b in buildings.iterrows():
                    design = design_map.get(b.code_era_proxy)
                    candidates = []
                    if design:
                        prefix = f"CR/{system}/{design}/H:"
                        candidates = sorted({int(k.split("/H:")[1].split("/")[0]) for k in model if k.startswith(prefix)})
                    height, height_delta = choose_height(float(b.storeys_reported), candidates, height_policy)
                    taxonomy = f"CR/{system}/{design}/H:{height}/RES" if design and height is not None else None
                    status = "eligible"
                    if b.code_era_proxy == "pre1975": status = "pre1975_design_unmapped"
                    elif design is None: status = "design_unmapped"
                    elif height is None: status = "height_unmapped"
                    elif taxonomy not in model: status = "taxonomy_absent"
                    for scope, event_keys in EVENT_SCOPES.items():
                        row = {"building_id": b.building_id, "scenario": scenario, "system_scenario": system,
                               "design_scenario": design_name, "height_policy": height_policy,
                               "event_scope": scope, "taxonomy": taxonomy, "height_class": height,
                               "height_delta_storeys": height_delta, "observed_ds": b.observed_ds,
                               "structural_damage_code": b.structural_damage_code,
                               "observed_loss_proxy": b.observed_loss_proxy, "status": status,
                               "source_supplemented_count": b.source_supplemented_count,
                               "soft_story_flag": b.soft_story_flag,
                               "age_years_reported": b.age_years_reported,
                               "storeys_reported": b.storeys_reported,
                               "vs30_source_mps": b.vs30_source_mps,
                               "min_mainshock_trace_distance_km": b.min_mainshock_trace_distance_km}
                        if status != "eligible":
                            results.append(row); continue
                        fun = model[taxonomy]
                        base = taxonomy.rsplit("/", 1)[0]
                        imt = fun["imt"]
                        event_im = np.array([im_lookup.loc[b.building_id, (e, imt)]
                                             if (e, imt) in im_lookup.columns else np.nan for e in event_keys], float)
                        row["imt"] = imt
                        row["n_events_required"] = len(event_keys)
                        if not np.isfinite(event_im).all():
                            row["status"] = "im_missing"
                            results.append(row); continue
                        row["max_im_g"] = float(event_im.max())
                        max_event_index = int(np.argmax(event_im))
                        row["max_event_key"] = event_keys[max_event_index]
                        row["max_event_label"] = EVENT_LABELS[event_keys[max_event_index]]
                        mean_max, _, outside = evaluate(fun, [event_im.max()], outside="nan")
                        event_losses = np.array([evaluate(fun, [x], outside="nan")[0][0] for x in event_im])
                        row["predicted_lr_max"] = mean_max[0]
                        row["predicted_lr_sequence_proxy"] = 1 - np.prod(1-event_losses) if np.isfinite(event_losses).all() else np.nan
                        row["outside_model_domain"] = bool(outside[0] or not np.isfinite(event_losses).all())
                        fp = fragility.get((base, imt))
                        if fp is None:
                            row["status"] = "fragility_missing"
                            results.append(row); continue
                        exc_events, adj_events = fragility_exceedance(fp, event_im)
                        exc_max_matrix, adj_max = fragility_exceedance(fp, [event_im.max()])
                        exc_max = exc_max_matrix[0]
                        exc_union = 1 - np.prod(1-exc_events, axis=0)
                        row["fragility_nesting_adjustment_max"] = float(max(adj_events.max(), adj_max[0]))
                        pmax = state_probabilities(exc_max[None, :])[0]
                        punion = state_probabilities(exc_union[None, :])[0]
                        for i in range(5):
                            row[f"fragility_max_p{i}"] = pmax[i]
                            row[f"fragility_union_proxy_p{i}"] = punion[i]
                        results.append(row)
    result = pd.DataFrame(results)
    write_csv(result, OUT / "building_results.csv")
    write_csv(result.groupby(["scenario", "event_scope", "status"], dropna=False).size().reset_index(name="n"),
              OUT / "exclusion_flow.csv")
    write_csv(result[result.status.eq("eligible")].groupby(
        ["scenario", "event_scope", "max_event_label", "imt"], dropna=False
    ).size().reset_index(name="n"), OUT / "max_event_contributions.csv")

    metrics_rows, calibration_rows = [], []
    eligible = result[result.status.eq("eligible")].copy()
    damage_mappings = {"code_C_as_DS4": lambda y: y,
                       "code_C_as_DS3": lambda y: np.minimum(y, 3)}
    for (scenario, scope), g in eligible.groupby(["scenario", "event_scope"]):
        for damage_mapping, transform in damage_mappings.items():
            y = transform(g.observed_ds.to_numpy(int))
            y_loss = pd.Series(y).map(DLR).to_numpy(float)
            for method, col in (("vulnerability_max", "predicted_lr_max"),
                                ("vulnerability_sequence_proxy", "predicted_lr_sequence_proxy")):
                metrics_rows.append({"scenario": scenario, "event_scope": scope,
                                     "damage_mapping": damage_mapping, "method": method,
                                     **vulnerability_metrics(y_loss, g[col])})
            for method, prefix in (("fragility_max", "fragility_max_p"),
                                   ("fragility_union_proxy", "fragility_union_proxy_p")):
                p = g[[f"{prefix}{i}" for i in range(5)]].to_numpy(float)
                scores = score_states(p, y)
                scores["accuracy_argmax"] = float((p.argmax(axis=1) == y).mean())
                metrics_rows.append({"scenario": scenario, "event_scope": scope,
                                     "damage_mapping": damage_mapping, "method": method, **scores})
                exceed = 1 - np.cumsum(p, axis=1)[:, :4]
                for ds in range(1, 5):
                    observed = (y >= ds).astype(float)
                    calibration_rows.append({"scenario": scenario, "event_scope": scope,
                        "damage_mapping": damage_mapping, "method": method,
                        "threshold": f"DS{ds}+", "n": len(g), "predicted_rate": exceed[:, ds-1].mean(),
                        "observed_rate": observed.mean(), "brier_binary": np.mean((exceed[:, ds-1]-observed)**2)})
    metrics = pd.DataFrame(metrics_rows)
    calibration = pd.DataFrame(calibration_rows)
    write_csv(metrics, OUT / "metrics.csv")
    write_csv(calibration, OUT / "fragility_calibration.csv")
    write_calibration_factors(eligible, damage_mappings)
    write_residual_correlations(eligible, damage_mappings)
    im_bin_calibration = write_im_bin_calibration(eligible, damage_mappings)
    common_scope_metrics(eligible, damage_mappings)
    plot_results(buildings, metrics, calibration, im_bin_calibration)
    manifest = {
        "source_sha256": sha256(SOURCE),
        "model_sha": "5974372ac3f4a99f25d0649eb030fbe596f23b36",
        "events_chronological": [label for _, label in EVENTS],
        "primary_candidate_scenario": "lfinf__baseline__exact",
        "warnings": [
            "Structural system and code/design class are sensitivity assumptions, not observed taxonomy fields.",
            "The meaning of source code C is unresolved; C-to-DS4 and C-to-DS3 mappings are separate sensitivities.",
            "Cumulative damage is compared using max-IM and explicitly uncalibrated union/remaining-value proxies.",
            "Observed loss is a damage-to-loss proxy, not repair-cost data.",
            "Fragility beta already contains record-to-record, building-to-building and DS-threshold components.",
            "SA(0.6) for events 3 and 4 is log-period/log-SA interpolation; event 4 covers 133 of 242 sites.",
            "Negative log score uses a 1e-15 probability floor for numerical tail truncation.",
        ],
    }
    (OUT / "run_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")


def write_calibration_factors(eligible: pd.DataFrame, damage_mappings) -> None:
    """Aggregate multiplicative factor, conditional on each explicit scenario.

    The interval is an ordinary building bootstrap.  It does not account for
    survey selection, spatial dependence, taxonomy ambiguity or IM uncertainty.
    """
    rng = np.random.default_rng(20230928)
    rows = []
    for (scenario, scope), g in eligible.groupby(["scenario", "event_scope"]):
        pred = g.predicted_lr_max.to_numpy(float)
        for damage_mapping, transform in damage_mappings.items():
            observed = pd.Series(transform(g.observed_ds.to_numpy(int))).map(DLR).to_numpy(float)
            ok = np.isfinite(observed) & np.isfinite(pred)
            observed_ok, pred_ok = observed[ok], pred[ok]
            factor = observed_ok.sum() / pred_ok.sum() if pred_ok.sum() > 0 else np.nan
            boot = []
            if len(observed_ok):
                for _ in range(2000):
                    idx = rng.integers(0, len(observed_ok), len(observed_ok))
                    denom = pred_ok[idx].sum()
                    if denom > 0:
                        boot.append(observed_ok[idx].sum() / denom)
            lo, hi = (np.quantile(boot, [0.025, 0.975]) if boot else (np.nan, np.nan))
            rows.append({"scenario": scenario, "event_scope": scope,
                         "damage_mapping": damage_mapping, "n": len(observed_ok),
                         "observed_over_predicted_factor": factor,
                         "bootstrap_ci_low": lo, "bootstrap_ci_high": hi,
                         "bootstrap_replicates": len(boot),
                         "interpretation": "conditional sensitivity result; not a national correction"})
    write_csv(pd.DataFrame(rows), OUT / "calibration_factors.csv")


def write_residual_correlations(eligible: pd.DataFrame, damage_mappings) -> None:
    rows = []
    covariates = ["vs30_source_mps", "min_mainshock_trace_distance_km",
                  "age_years_reported", "storeys_reported"]
    for (scenario, scope), g in eligible.groupby(["scenario", "event_scope"]):
        pred = g.predicted_lr_max.to_numpy(float)
        for damage_mapping, transform in damage_mappings.items():
            observed = pd.Series(transform(g.observed_ds.to_numpy(int))).map(DLR).to_numpy(float)
            residual = observed - pred
            for covariate in covariates:
                x = g[covariate].to_numpy(float)
                ok = np.isfinite(x) & np.isfinite(residual)
                result = spearmanr(x[ok], residual[ok]) if ok.sum() > 2 else None
                rows.append({"scenario": scenario, "event_scope": scope,
                             "damage_mapping": damage_mapping, "residual": "observed_minus_predicted",
                             "covariate": covariate, "n": int(ok.sum()),
                             "spearman_rho": result.statistic if result else np.nan,
                             "p_value_unadjusted": result.pvalue if result else np.nan})
    write_csv(pd.DataFrame(rows), OUT / "residual_correlations.csv")


def write_im_bin_calibration(eligible: pd.DataFrame, damage_mappings) -> pd.DataFrame:
    """Equal-count IM bins for loss-proxy and DS3+ calibration displays."""
    rows = []
    for (scenario, scope, imt), g0 in eligible.groupby(["scenario", "event_scope", "imt"]):
        g0 = g0[np.isfinite(g0.max_im_g) & np.isfinite(g0.predicted_lr_max)].copy()
        if len(g0) < 20:
            continue
        g0["im_bin"] = pd.qcut(g0.max_im_g, q=min(4, len(g0)), duplicates="drop")
        for damage_mapping, transform in damage_mappings.items():
            g = g0.copy()
            g["observed_ds_mapped"] = transform(g.observed_ds.to_numpy(int))
            g["observed_loss_mapped"] = g.observed_ds_mapped.map(DLR)
            for interval, z in g.groupby("im_bin", observed=True):
                rows.append({"scenario": scenario, "event_scope": scope,
                             "damage_mapping": damage_mapping, "imt": imt,
                             "im_bin_left_g": interval.left, "im_bin_right_g": interval.right,
                             "n": len(z), "mean_im_g": z.max_im_g.mean(),
                             "predicted_loss_mean": z.predicted_lr_max.mean(),
                             "observed_loss_proxy_mean": z.observed_loss_mapped.mean(),
                             "predicted_ds3plus_mean": z.fragility_max_p3.add(z.fragility_max_p4).mean(),
                             "observed_ds3plus_rate": (z.observed_ds_mapped >= 3).mean()})
    result = pd.DataFrame(rows)
    write_csv(result, OUT / "im_bin_calibration.csv")
    return result


def common_scope_metrics(eligible: pd.DataFrame, damage_mappings) -> None:
    """Compare two- and four-event representations on identical buildings."""
    rows = []
    for scenario, g in eligible.groupby("scenario"):
        ids_by_scope = {scope: set(z.building_id) for scope, z in g.groupby("event_scope")}
        if set(EVENT_SCOPES) - set(ids_by_scope):
            continue
        common = set.intersection(*(ids_by_scope[s] for s in EVENT_SCOPES))
        for scope in EVENT_SCOPES:
            z = g[(g.event_scope == scope) & g.building_id.isin(common)].copy()
            for damage_mapping, transform in damage_mappings.items():
                y = transform(z.observed_ds.to_numpy(int))
                y_loss = pd.Series(y).map(DLR).to_numpy(float)
                rows.append({"scenario": scenario, "event_scope": scope,
                             "damage_mapping": damage_mapping, "method": "vulnerability_max",
                             "common_n": len(common), **vulnerability_metrics(y_loss, z.predicted_lr_max)})
                p = z[[f"fragility_max_p{i}" for i in range(5)]].to_numpy(float)
                scores = score_states(p, y)
                scores["accuracy_argmax"] = float((p.argmax(axis=1) == y).mean())
                rows.append({"scenario": scenario, "event_scope": scope,
                             "damage_mapping": damage_mapping, "method": "fragility_max",
                             "common_n": len(common), **scores})
    write_csv(pd.DataFrame(rows), OUT / "common_event_scope_metrics.csv")


def plot_results(buildings: pd.DataFrame, metrics: pd.DataFrame, calibration: pd.DataFrame,
                 im_bin_calibration: pd.DataFrame) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    counts = buildings.structural_damage_code.value_counts().reindex(["N", "L", "M", "S", "C"], fill_value=0)
    fig, ax = plt.subplots(figsize=(6.5, 4.2))
    counts.plot.bar(ax=ax, color="#3c6e71", rot=0)
    ax.set(xlabel="ACI structural damage code", ylabel="Buildings", title="ACI133 analytic sample (n=242)")
    fig.tight_layout(); fig.savefig(OUT / "damage_distribution.png", dpi=200); plt.close(fig)

    c = calibration[(calibration.scenario == "lfinf__baseline__exact") &
                    (calibration.event_scope == "four_survey_preceding_events") &
                    (calibration.damage_mapping == "code_C_as_DS4") &
                    (calibration.method == "fragility_max")]
    if len(c):
        fig, ax = plt.subplots(figsize=(5.4, 5.0))
        ax.plot([0, 1], [0, 1], color="0.55", lw=1, ls="--")
        ax.scatter(c.predicted_rate, c.observed_rate, s=60, color="#284b63")
        for _, row in c.iterrows():
            ax.annotate(row.threshold, (row.predicted_rate, row.observed_rate), xytext=(5, 4),
                        textcoords="offset points", fontsize=9)
        ax.set(xlim=(0,1), ylim=(0,1), xlabel="Mean predicted exceedance probability",
               ylabel="Observed exceedance frequency", title="Fragility calibration: baseline exact-height subset")
        fig.tight_layout(); fig.savefig(OUT / "fragility_calibration_baseline.png", dpi=200); plt.close(fig)

    m = metrics[(metrics.method == "vulnerability_max") &
                (metrics.event_scope == "four_survey_preceding_events") &
                (metrics.damage_mapping == "code_C_as_DS4")].copy()
    if len(m):
        m = m.sort_values("bias_pred_obs")
        fig, ax = plt.subplots(figsize=(8, max(4, len(m)*0.22)))
        bars = ax.barh(m.scenario, m.bias_pred_obs, color="#6b8e23")
        for bar, n in zip(bars, m.n):
            ax.text(bar.get_width() + 0.015, bar.get_y() + bar.get_height()/2,
                    f"n={int(n)}", va="center", fontsize=7)
        ax.axvline(1, color="0.25", lw=1, ls="--")
        ax.set(xlabel="Predicted / observed damage-derived loss", ylabel="Taxonomy scenario",
               title="Taxonomy sensitivity (conditional available samples)")
        fig.tight_layout(); fig.savefig(OUT / "vulnerability_bias_scenarios.png", dpi=200); plt.close(fig)

    b = im_bin_calibration[(im_bin_calibration.scenario == "lfinf__baseline__exact") &
                           (im_bin_calibration.event_scope == "four_survey_preceding_events") &
                           (im_bin_calibration.damage_mapping == "code_C_as_DS4")]
    if len(b):
        fig, ax = plt.subplots(figsize=(6.3, 4.5))
        ax.plot(b.mean_im_g, b.predicted_loss_mean, marker="o", label="GEM predicted mean loss ratio")
        ax.plot(b.mean_im_g, b.observed_loss_proxy_mean, marker="s", label="Observed damage-loss proxy")
        ax.set(xlabel=f"{b.imt.iloc[0]} (g), equal-count bins", ylabel="Mean loss ratio",
               title="Intensity-bin calibration: baseline exact-height subset")
        ax.legend(frameon=False)
        fig.tight_layout(); fig.savefig(OUT / "im_bin_calibration_baseline.png", dpi=200); plt.close(fig)


def main() -> None:
    AUDIT.mkdir(parents=True, exist_ok=True)
    OUT.mkdir(parents=True, exist_ok=True)
    buildings = prepare()
    ims = assign_ims(buildings)
    analyse(buildings, ims)
    print(f"ACI133 analysis complete: {len(buildings)} buildings")


if __name__ == "__main__":
    main()
